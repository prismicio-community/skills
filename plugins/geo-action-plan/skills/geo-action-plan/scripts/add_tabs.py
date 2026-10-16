#!/usr/bin/env python3
"""Add the Cited Sources and Topic Clusters tabs to a workbook built by build_geo_workbook.py.

Usage:
    python add_tabs.py run.json output.xlsx
"""
import json
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

run = json.load(open(sys.argv[1]))
wb = load_workbook(sys.argv[2])
H = PatternFill("solid", fgColor="1F3864")
HF = Font(color="FFFFFF", bold=True)


def sheet(name, head, rows, widths, after):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name, wb.sheetnames.index(after) + 1)
    ws.append(head)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for c in ws[1]:
        c.fill, c.font = H, HF
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"


cid = {}
for c in run.get("clusters", []):
    for m in c["members"]:
        cid.setdefault(m["url"].rstrip("/").lower(), c["id"])

src = sorted(run.get("citation_sources", []), key=lambda s: (-(s.get("retrievals") or 0), -(s.get("citations") or 0)))
sheet(
    "Cited Sources",
    ["URL", "Title", "Domain", "Page type", "Retrievals", "Citations", "Citation rate", "Brands mentioned", "Cluster"],
    [
        [s["url"], s.get("title", ""), s.get("domain", ""), s.get("classification", ""), s.get("retrievals"),
         s.get("citations"), s.get("citation_rate"), ", ".join(s.get("mentioned_brands", [])),
         cid.get(s["url"].rstrip("/").lower(), "")]
        for s in src
    ],
    [60, 50, 24, 14, 11, 11, 11, 50, 9],
    "Read Me",
)

rows = []
for c in run.get("clusters", []):
    for i, m in enumerate(c["members"]):
        first = i == 0
        rows.append([
            c["id"] if first else "", c["topic"] if first else "", c.get("fit", "") if first else "",
            c.get("action", "") if first else "", c.get("existing_url", "") if first else "",
            c.get("plan_row", "") if first else "", m["title"], m["url"], m["layer"], m["metric"],
        ])
sheet(
    "Topic Clusters",
    ["Cluster", "Topic (your title)", "Fit", "Action", "Existing page on your site", "Plan row",
     "Member page title", "Member URL", "Layer", "Metric"],
    rows,
    [8, 46, 8, 10, 40, 8, 50, 55, 9, 26],
    "Cited Sources",
)
wb.save(sys.argv[2])
print("tabs:", wb.sheetnames)
